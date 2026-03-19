from docx import Document

from app.models.strategy_lab import LocatorType
from app.utils.file_parser import FileParser


def test_extract_document_supports_docx_with_locators(tmp_path):
    docx_path = tmp_path / "strategy.docx"

    document = Document()
    document.add_heading("Strategic Outlook", level=1)
    document.add_paragraph("Defense modernization is a near-term opportunity.")
    document.add_paragraph("Healthcare pilots require a different buying path.")
    document.save(docx_path)

    extracted = FileParser.extract_document(str(docx_path))

    assert extracted.extension == "docx"
    assert "Strategic Outlook" in extracted.text
    assert extracted.segments[0].locator_type == LocatorType.HEADING
    assert extracted.segments[0].locator == "heading:1"
    assert extracted.segments[1].locator == "paragraph:2"
    assert FileParser.extract_text(str(docx_path)) == extracted.text


def test_extract_document_handles_docx_paragraphs_without_style(monkeypatch, tmp_path):
    docx_path = tmp_path / "styleless.docx"
    docx_path.write_bytes(b"fake")

    class FakeParagraph:
        def __init__(self, text, style):
            self.text = text
            self.style = style

    class FakeDocument:
        paragraphs = [
            FakeParagraph("Untitled section", None),
            FakeParagraph("Market proof still loads.", None),
        ]

    monkeypatch.setattr("docx.Document", lambda _: FakeDocument())

    extracted = FileParser.extract_document(str(docx_path))

    assert extracted.extension == "docx"
    assert extracted.segments[0].locator_type == LocatorType.PARAGRAPH
    assert extracted.segments[0].locator == "paragraph:1"
