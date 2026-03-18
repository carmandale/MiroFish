"""
文件解析工具
支持 PDF、DOCX、Markdown、TXT 文件的文本提取
"""

import os
import uuid
from pathlib import Path
from typing import List

from ..models.strategy_lab import (
    DocumentSegment,
    ExtractedDocument,
    LocatorType,
)


def _read_text_with_fallback(file_path: str) -> str:
    """
    读取文本文件，UTF-8失败时自动探测编码。
    
    采用多级回退策略：
    1. 首先尝试 UTF-8 解码
    2. 使用 charset_normalizer 检测编码
    3. 回退到 chardet 检测编码
    4. 最终使用 UTF-8 + errors='replace' 兜底
    
    Args:
        file_path: 文件路径
        
    Returns:
        解码后的文本内容
    """
    data = Path(file_path).read_bytes()
    
    # 首先尝试 UTF-8
    try:
        return data.decode('utf-8')
    except UnicodeDecodeError:
        pass
    
    # 尝试使用 charset_normalizer 检测编码
    encoding = None
    try:
        from charset_normalizer import from_bytes
        best = from_bytes(data).best()
        if best and best.encoding:
            encoding = best.encoding
    except Exception:
        pass
    
    # 回退到 chardet
    if not encoding:
        try:
            import chardet
            result = chardet.detect(data)
            encoding = result.get('encoding') if result else None
        except Exception:
            pass
    
    # 最终兜底：使用 UTF-8 + replace
    if not encoding:
        encoding = 'utf-8'
    
    return data.decode(encoding, errors='replace')


class FileParser:
    """文件解析器"""
    
    SUPPORTED_EXTENSIONS = {'.pdf', '.docx', '.md', '.markdown', '.txt'}
    
    @classmethod
    def extract_text(cls, file_path: str) -> str:
        """
        从文件中提取文本
        
        Args:
            file_path: 文件路径
            
        Returns:
            提取的文本内容
        """
        path = Path(file_path)
        
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")
        
        return cls.extract_document(file_path).text

    @classmethod
    def extract_document(
        cls,
        file_path: str,
        *,
        original_filename: str | None = None,
        saved_filename: str | None = None,
    ) -> ExtractedDocument:
        """
        从文件中提取文本和定位信息
        """
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        suffix = path.suffix.lower()
        if suffix not in cls.SUPPORTED_EXTENSIONS:
            raise ValueError(f"不支持的文件格式: {suffix}")

        if suffix == '.pdf':
            text, segments = cls._extract_pdf_document(file_path)
        elif suffix == '.docx':
            text, segments = cls._extract_docx_document(file_path)
        elif suffix in {'.md', '.markdown'}:
            text, segments = cls._extract_text_document(
                file_path,
                extension=suffix,
            )
        elif suffix == '.txt':
            text, segments = cls._extract_text_document(
                file_path,
                extension=suffix,
            )
        else:
            raise ValueError(f"无法处理的文件格式: {suffix}")

        return ExtractedDocument(
            document_id=f"doc_{uuid.uuid4().hex[:12]}",
            original_filename=original_filename or path.name,
            saved_filename=saved_filename or path.name,
            file_path=str(path),
            extension=suffix.lstrip('.'),
            text=text,
            segments=segments,
        )
    
    @staticmethod
    def _extract_pdf_document(file_path: str) -> tuple[str, List[DocumentSegment]]:
        """从 PDF 提取文本和页级定位信息"""
        try:
            import fitz  # PyMuPDF
        except ImportError:
            raise ImportError("需要安装PyMuPDF: pip install PyMuPDF")
        
        text_parts: List[str] = []
        segments: List[DocumentSegment] = []
        cursor = 0
        with fitz.open(file_path) as doc:
            for index, page in enumerate(doc, start=1):
                text = page.get_text()
                normalized = text.strip()
                if not normalized:
                    continue

                if text_parts:
                    cursor += 2
                start_char = cursor
                text_parts.append(normalized)
                cursor += len(normalized)
                segments.append(
                    DocumentSegment(
                        locator_type=LocatorType.PAGE,
                        locator=f"page:{index}",
                        start_char=start_char,
                        end_char=cursor,
                        text=normalized,
                    )
                )

        return "\n\n".join(text_parts), segments
    
    @staticmethod
    def _extract_docx_document(file_path: str) -> tuple[str, List[DocumentSegment]]:
        """从 DOCX 提取文本和段落级定位信息"""
        try:
            from docx import Document
        except ImportError:
            raise ImportError("需要安装 python-docx: pip install python-docx")

        doc = Document(file_path)
        blocks = []
        for index, paragraph in enumerate(doc.paragraphs, start=1):
            content = paragraph.text.strip()
            if not content:
                continue

            style_name = (paragraph.style.name or "").lower()
            if "heading" in style_name:
                locator_type = LocatorType.HEADING
                locator = f"heading:{index}"
            else:
                locator_type = LocatorType.PARAGRAPH
                locator = f"paragraph:{index}"

            blocks.append((locator_type, locator, content))

        return FileParser._build_segmented_document(blocks)
    
    @staticmethod
    def _extract_text_document(
        file_path: str,
        *,
        extension: str,
    ) -> tuple[str, List[DocumentSegment]]:
        """从纯文本或 Markdown 提取文本和段落定位信息"""
        raw_text = _read_text_with_fallback(file_path)
        blocks = []
        paragraph_index = 0

        for block in raw_text.replace('\r\n', '\n').replace('\r', '\n').split('\n\n'):
            content = block.strip()
            if not content:
                continue

            paragraph_index += 1
            locator_type = LocatorType.SECTION if extension in {'.md', '.markdown'} and content.startswith('#') else LocatorType.PARAGRAPH
            prefix = "section" if locator_type == LocatorType.SECTION else "paragraph"
            blocks.append((locator_type, f"{prefix}:{paragraph_index}", content))

        return FileParser._build_segmented_document(blocks)

    @staticmethod
    def _build_segmented_document(
        blocks: List[tuple[LocatorType, str, str]],
    ) -> tuple[str, List[DocumentSegment]]:
        text_parts: List[str] = []
        segments: List[DocumentSegment] = []
        cursor = 0

        for locator_type, locator, content in blocks:
            if text_parts:
                cursor += 2
            start_char = cursor
            text_parts.append(content)
            cursor += len(content)
            segments.append(
                DocumentSegment(
                    locator_type=locator_type,
                    locator=locator,
                    start_char=start_char,
                    end_char=cursor,
                    text=content,
                )
            )

        return "\n\n".join(text_parts), segments
    
    @classmethod
    def extract_from_multiple(cls, file_paths: List[str]) -> str:
        """
        从多个文件提取文本并合并
        
        Args:
            file_paths: 文件路径列表
            
        Returns:
            合并后的文本
        """
        all_texts = []
        
        for i, file_path in enumerate(file_paths, 1):
            try:
                text = cls.extract_document(file_path).text
                filename = Path(file_path).name
                all_texts.append(f"=== 文档 {i}: {filename} ===\n{text}")
            except Exception as e:
                all_texts.append(f"=== 文档 {i}: {file_path} (提取失败: {str(e)}) ===")
        
        return "\n\n".join(all_texts)


def split_text_into_chunks(
    text: str, 
    chunk_size: int = 500, 
    overlap: int = 50
) -> List[str]:
    """
    将文本分割成小块
    
    Args:
        text: 原始文本
        chunk_size: 每块的字符数
        overlap: 重叠字符数
        
    Returns:
        文本块列表
    """
    if len(text) <= chunk_size:
        return [text] if text.strip() else []
    
    chunks = []
    start = 0
    
    while start < len(text):
        end = start + chunk_size
        
        # 尝试在句子边界处分割
        if end < len(text):
            # 查找最近的句子结束符
            for sep in ['。', '！', '？', '.\n', '!\n', '?\n', '\n\n', '. ', '! ', '? ']:
                last_sep = text[start:end].rfind(sep)
                if last_sep != -1 and last_sep > chunk_size * 0.3:
                    end = start + last_sep + len(sep)
                    break
        
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        
        # 下一个块从重叠位置开始
        start = end - overlap if end < len(text) else len(text)
    
    return chunks
